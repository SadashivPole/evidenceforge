"""API tests for workspace-scoped questionnaire XLSX import."""

from __future__ import annotations

from io import BytesIO
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import WorkspaceMembership, WorkspaceRole
from app.questionnaires.persistence.errors import (
    QuestionnaireNotFoundError,
    QuestionnairePersistenceError,
    QuestionnairePersistenceValidationError,
)
from app.questionnaires.persistence.models import (
    QuestionnaireImportAttempt,
    QuestionnaireVersion,
)
from tests.conftest import (
    auth_headers,
    create_principal,
    create_workspace_with_owner,
)


def make_workbook_bytes(
    *,
    instruction: str = "This worksheet contains guidance only.",
) -> bytes:
    """Create a deterministic questionnaire workbook for API tests."""

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Security"

    worksheet.append(["Section", "Question"])
    worksheet.append(
        [
            "Access Control",
            "Do you enforce multi-factor authentication?",
        ]
    )
    worksheet.append(
        [
            None,
            "Are privileged accounts reviewed regularly?",
        ]
    )

    instructions = workbook.create_sheet("Instructions")
    instructions.append([instruction])

    buffer = BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def make_invalid_source_id_workbook_bytes() -> bytes:
    """Create a workbook with an over-length explicit source question ID."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Security"
    worksheet.append(["Question ID", "Question"])
    worksheet.append(["Q" * 256, "Do you use MFA?"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _post_import(
    client: TestClient,
    principal,
    workspace,
    workbook: bytes,
    *,
    filename: str = "questionnaire.xlsx",
):
    return client.post(
        f"/workspaces/{workspace.id}/questionnaires/import",
        headers=auth_headers(principal),
        files={
            "file": (
                filename,
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_questionnaire_xlsx_import_preview(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="owner@example.com",
        display_name="Owner",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Questionnaire Workspace",
    )

    response = _post_import(
        client,
        principal,
        workspace,
        make_workbook_bytes(),
        filename="vendor-questionnaire.xlsx",
    )

    assert response.status_code == 200

    body = response.json()

    assert body["outcome"] == "CREATED"
    assert body["version_number"] == 1
    assert body["version_id"]
    assert body["import_attempt_id"]
    assert body["name"] == "vendor-questionnaire"
    assert body["source_filename"] == "vendor-questionnaire.xlsx"
    assert body["status"] == "IMPORTED"
    assert body["parser_version"] == "xlsx-import-v1"
    assert body["normalization_version"] == "question-normalization-v1"
    assert body["question_identity_version"] == "question-identity-v2"
    assert body["hash_version"] == "questionnaire-hash-v1"
    assert len(body["normalized_questionnaire_sha256"]) == 64

    assert len(body["questions"]) == 2

    first = body["questions"][0]

    assert first["ordinal"] == 1
    assert first["sheet_name"] == "Security"
    assert first["source_row"] == 2
    assert first["question_text"] == ("Do you enforce multi-factor authentication?")
    assert first["identity_kind"] == "fallback"
    assert first["source_question_id"] is None
    assert first["normalized_sheet_name"] == "Security"
    assert first["normalized_question_text"] == ("Do you enforce multi-factor authentication?")
    assert first["section_path"] == ["Access Control"]
    assert first["normalized_section_path"] == ["Access Control"]

    second = body["questions"][1]

    assert second["ordinal"] == 2
    assert second["source_row"] == 3
    assert second["section_path"] == ["Access Control"]

    assert len(body["imported_sheets"]) == 1
    assert body["imported_sheets"][0]["sheet_name"] == "Security"
    assert body["imported_sheets"][0]["question_count"] == 2

    assert len(body["ignored_sheets"]) == 1
    assert body["ignored_sheets"][0]["sheet_name"] == "Instructions"

    version = db_session.scalar(
        select(QuestionnaireVersion).where(
            QuestionnaireVersion.workspace_id == workspace.id,
        )
    )
    attempt = db_session.scalar(
        select(QuestionnaireImportAttempt).where(
            QuestionnaireImportAttempt.workspace_id == workspace.id,
        )
    )

    assert version is not None
    assert attempt is not None
    assert body["questionnaire_id"] == str(version.questionnaire_id)
    assert body["version_id"] == str(version.id)
    assert body["import_attempt_id"] == str(attempt.id)


def test_invalid_source_question_id_returns_422(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="invalid-source-id@example.com",
        display_name="Invalid Source ID",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Invalid Source ID Workspace",
    )

    response = _post_import(
        client,
        principal,
        workspace,
        make_invalid_source_id_workbook_bytes(),
    )

    assert response.status_code == 422


def test_question_ids_are_stable_across_repeated_api_imports(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="stable@example.com",
        display_name="Stable Import",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Stable Workspace",
    )

    workbook = make_workbook_bytes()

    first = _post_import(client, principal, workspace, workbook)
    second = _post_import(client, principal, workspace, workbook)

    assert first.status_code == 200
    assert second.status_code == 200

    first_body = first.json()
    second_body = second.json()

    assert first_body["outcome"] == "CREATED"
    assert second_body["outcome"] == "IDENTICAL_DUPLICATE"
    assert first_body["questionnaire_id"] == second_body["questionnaire_id"]
    assert first_body["version_id"] == second_body["version_id"]
    assert first_body["version_number"] == second_body["version_number"] == 1
    assert first_body["import_attempt_id"] != second_body["import_attempt_id"]

    first_ids = [question["question_id"] for question in first_body["questions"]]
    second_ids = [question["question_id"] for question in second_body["questions"]]

    assert first_ids == second_ids
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(QuestionnaireVersion)
            .where(QuestionnaireVersion.workspace_id == workspace.id)
        )
        == 1
    )
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(QuestionnaireImportAttempt)
            .where(QuestionnaireImportAttempt.workspace_id == workspace.id)
        )
        == 2
    )


def test_canonical_duplicate_returns_existing_version(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="canonical@example.com",
        display_name="Canonical Duplicate",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Canonical Workspace",
    )

    first = _post_import(
        client,
        principal,
        workspace,
        make_workbook_bytes(instruction="First guidance"),
    )
    second = _post_import(
        client,
        principal,
        workspace,
        make_workbook_bytes(instruction="Different guidance"),
    )

    assert first.status_code == 200
    assert second.status_code == 200

    first_body = first.json()
    second_body = second.json()

    assert first_body["outcome"] == "CREATED"
    assert second_body["outcome"] == "CANONICAL_DUPLICATE"
    assert first_body["questionnaire_id"] == second_body["questionnaire_id"]
    assert first_body["version_id"] == second_body["version_id"]
    assert first_body["version_number"] == second_body["version_number"] == 1
    assert first_body["import_attempt_id"] != second_body["import_attempt_id"]
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(QuestionnaireVersion)
            .where(QuestionnaireVersion.workspace_id == workspace.id)
        )
        == 1
    )
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(QuestionnaireImportAttempt)
            .where(QuestionnaireImportAttempt.workspace_id == workspace.id)
        )
        == 2
    )


@pytest.mark.parametrize(
    ("error_type", "expected_status", "expected_detail"),
    [
        (
            QuestionnairePersistenceValidationError,
            422,
            "Questionnaire import failed persistence validation",
        ),
        (
            QuestionnaireNotFoundError,
            404,
            "Questionnaire could not be resolved",
        ),
        (
            QuestionnairePersistenceError,
            500,
            "Questionnaire persistence failed",
        ),
    ],
)
def test_persistence_failures_return_safe_http_errors(
    client: TestClient,
    db_session: Session,
    error_type: type[QuestionnairePersistenceError],
    expected_status: int,
    expected_detail: str,
) -> None:
    principal = create_principal(
        db_session,
        email=f"persistence-{expected_status}@example.com",
        display_name="Persistence Failure",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name=f"Persistence Failure {expected_status}",
    )

    with patch(
        "app.api.routes.questionnaires.persist_import",
        side_effect=error_type("test failure"),
    ):
        response = _post_import(
            client,
            principal,
            workspace,
            make_workbook_bytes(),
        )

    assert response.status_code == expected_status
    assert response.json() == {"detail": expected_detail}
    assert "test failure" not in response.text


def test_unsupported_questionnaire_extension_returns_415(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="extension@example.com",
        display_name="Extension",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Extension Workspace",
    )

    response = _post_import(
        client,
        principal,
        workspace,
        b"not an xlsx workbook",
        filename="questionnaire.xls",
    )

    assert response.status_code == 415


def test_corrupt_questionnaire_returns_422(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="corrupt@example.com",
        display_name="Corrupt",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Corrupt Workspace",
    )

    response = _post_import(
        client,
        principal,
        workspace,
        b"definitely not a real workbook",
        filename="corrupt.xlsx",
    )

    assert response.status_code == 422


def test_oversized_questionnaire_returns_413(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="large@example.com",
        display_name="Large File",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Large File Workspace",
    )

    oversized = b"x" * (10 * 1024 * 1024 + 1)

    response = _post_import(
        client,
        principal,
        workspace,
        oversized,
        filename="large.xlsx",
    )

    assert response.status_code == 413


def test_viewer_cannot_import_questionnaire(
    client: TestClient,
    db_session: Session,
) -> None:
    owner = create_principal(
        db_session,
        email="owner-questionnaire@example.com",
        display_name="Owner",
    )
    workspace = create_workspace_with_owner(
        db_session,
        owner,
        name="Viewer Workspace",
    )

    viewer = create_principal(
        db_session,
        email="viewer-questionnaire@example.com",
        display_name="Viewer",
    )

    db_session.add(
        WorkspaceMembership(
            workspace_id=workspace.id,
            user_id=viewer.user.id,
            role=WorkspaceRole.VIEWER,
        )
    )
    db_session.commit()

    response = _post_import(
        client,
        viewer,
        workspace,
        make_workbook_bytes(),
    )

    assert response.status_code == 403


def test_cross_workspace_questionnaire_import_returns_404(
    client: TestClient,
    db_session: Session,
) -> None:
    owner = create_principal(
        db_session,
        email="cross-owner@example.com",
        display_name="Owner",
    )
    workspace = create_workspace_with_owner(
        db_session,
        owner,
        name="Private Workspace",
    )

    other = create_principal(
        db_session,
        email="cross-other@example.com",
        display_name="Other User",
    )

    response = _post_import(
        client,
        other,
        workspace,
        make_workbook_bytes(),
    )

    assert response.status_code == 404


def test_unauthenticated_questionnaire_import_returns_401(
    client: TestClient,
    db_session: Session,
) -> None:
    principal = create_principal(
        db_session,
        email="unauth@example.com",
        display_name="Unauthenticated Test",
    )
    workspace = create_workspace_with_owner(
        db_session,
        principal,
        name="Auth Workspace",
    )

    response = client.post(
        f"/workspaces/{workspace.id}/questionnaires/import",
        files={
            "file": (
                "questionnaire.xlsx",
                make_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 401
