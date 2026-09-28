"""Deterministic XLSX questionnaire importer."""

from __future__ import annotations

from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.questionnaires.normalization import normalize_filename, normalize_text
from app.questionnaires.service import build_question, build_questionnaire
from app.questionnaires.xlsx.errors import (
    AmbiguousQuestionColumnError,
    InvalidQuestionValueError,
    InvalidXlsxWorkbookError,
    NoQuestionColumnError,
    TooManyQuestionsError,
    TooManySheetsError,
    UnsupportedXlsxExtensionError,
    XlsxFileTooLargeError,
)
from app.questionnaires.xlsx.policy import (
    QUESTION_HEADERS,
    SECTION_HEADERS,
    XLSX_IMPORT_VERSION,
    XlsxImportPolicy,
)
from app.questionnaires.xlsx.types import (
    IgnoredSheet,
    ImportedSheet,
    XlsxImportResult,
)


def normalize_header(value: object) -> str:
    """Normalize a workbook header for deterministic matching."""

    if value is None:
        return ""

    return normalize_text(str(value)).casefold()


def _cell_text(value: object) -> str | None:
    """Convert a worksheet cell to normalized text."""

    if value is None:
        return None

    normalized = normalize_text(str(value))

    if not normalized:
        return None

    return normalized


def _validate_filename(filename: str) -> None:
    """Validate the supplied filename extension."""

    suffix = PurePath(filename).suffix.casefold()

    if suffix != ".xlsx":
        raise UnsupportedXlsxExtensionError(
            f"unsupported questionnaire extension: {suffix or '<none>'}"
        )


def _find_header(
    *,
    worksheet: object,
    allowed_headers: frozenset[str],
    max_scan_rows: int,
) -> tuple[int, int, str] | None:
    """
    Locate one supported header.

    Returns:
        (row_number, column_number, original_header_text)
    """

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=max_scan_rows,
            values_only=True,
        ),
        start=1,
    ):
        for column_number, value in enumerate(row, start=1):
            normalized = normalize_header(value)

            if normalized in allowed_headers:
                original = _cell_text(value)

                if original is not None:
                    return row_number, column_number, original

    return None


def _find_headers(
    *,
    worksheet: object,
    allowed_headers: frozenset[str],
    max_scan_rows: int,
) -> list[tuple[int, int, str]]:
    """Return all matching headers in the scan window."""

    matches: list[tuple[int, int, str]] = []

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=max_scan_rows,
            values_only=True,
        ),
        start=1,
    ):
        for column_number, value in enumerate(row, start=1):
            normalized = normalize_header(value)

            if normalized in allowed_headers:
                original = _cell_text(value)

                if original is not None:
                    matches.append((row_number, column_number, original))

    return matches


def _find_question_header(
    *,
    worksheet: object,
    policy: XlsxImportPolicy,
) -> tuple[int, int, str] | None:
    """Find exactly one question header in a worksheet."""

    matches = _find_headers(
        worksheet=worksheet,
        allowed_headers=QUESTION_HEADERS,
        max_scan_rows=policy.max_header_scan_rows,
    )

    if not matches:
        return None

    rows = {match[0] for match in matches}

    earliest_row = min(rows)

    earliest_matches = [
        match
        for match in matches
        if match[0] == earliest_row
    ]

    if len(earliest_matches) > 1:
        raise AmbiguousQuestionColumnError(
            f"worksheet {worksheet.title!r} contains multiple question headers "
            f"on row {earliest_row}"
        )

    return earliest_matches[0]


def _find_section_header(
    *,
    worksheet: object,
    header_row: int,
    policy: XlsxImportPolicy,
) -> tuple[int, str] | None:
    """Find one optional section header on the question header row."""

    matches = [
        match
        for match in _find_headers(
            worksheet=worksheet,
            allowed_headers=SECTION_HEADERS,
            max_scan_rows=policy.max_header_scan_rows,
        )
        if match[0] == header_row
    ]

    if not matches:
        return None

    if len(matches) > 1:
        raise InvalidXlsxWorkbookError(
            f"worksheet {worksheet.title!r} contains multiple section headers "
            f"on question header row {header_row}"
        )

    _, column_number, original = matches[0]

    return column_number, original


def _questionnaire_name_from_filename(filename: str) -> str:
    """Derive a deterministic display name from the workbook filename."""

    stem = PurePath(filename).stem
    normalized = normalize_filename(stem)

    return normalized or "Imported Questionnaire"


def import_xlsx(
    *,
    data: bytes,
    filename: str,
    policy: XlsxImportPolicy | None = None,
    questionnaire_name: str | None = None,
) -> XlsxImportResult:
    """
    Parse an XLSX questionnaire deterministically.

    Only explicit supported question headers are interpreted.
    Worksheets without a supported question header are reported as ignored.
    """

    active_policy = policy or XlsxImportPolicy()

    _validate_filename(filename)

    if len(data) > active_policy.max_file_size_bytes:
        raise XlsxFileTooLargeError(
            f"questionnaire exceeds maximum size of "
            f"{active_policy.max_file_size_bytes} bytes"
        )

    if not data:
        raise InvalidXlsxWorkbookError("questionnaire file is empty")

    try:
        workbook = load_workbook(
            filename=BytesIO(data),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as exc:
        raise InvalidXlsxWorkbookError(
            "questionnaire is not a valid XLSX workbook"
        ) from exc

    try:
        sheet_names = tuple(workbook.sheetnames)

        if len(sheet_names) > active_policy.max_sheets:
            raise TooManySheetsError(
                f"workbook contains {len(sheet_names)} sheets; "
                f"maximum is {active_policy.max_sheets}"
            )

        questions = []
        imported_sheets: list[ImportedSheet] = []
        ignored_sheets: list[IgnoredSheet] = []

        for worksheet in workbook.worksheets:
            question_header = _find_question_header(
                worksheet=worksheet,
                policy=active_policy,
            )

            if question_header is None:
                ignored_sheets.append(
                    IgnoredSheet(
                        sheet_name=worksheet.title,
                        reason="no supported question header found",
                    )
                )
                continue

            header_row, question_column, question_header_text = question_header

            section_header = _find_section_header(
                worksheet=worksheet,
                header_row=header_row,
                policy=active_policy,
            )

            section_column: int | None = None
            section_header_text: str | None = None

            if section_header is not None:
                section_column, section_header_text = section_header

            current_section: tuple[str, ...] = ()
            sheet_question_count = 0

            for source_row, row in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    values_only=True,
                ),
                start=header_row + 1,
            ):
                question_value = (
                    row[question_column - 1]
                    if question_column <= len(row)
                    else None
                )

                if section_column is not None:
                    section_value = (
                        row[section_column - 1]
                        if section_column <= len(row)
                        else None
                    )

                    normalized_section = _cell_text(section_value)

                    if normalized_section is not None:
                        current_section = (normalized_section,)

                normalized_question = _cell_text(question_value)

                if normalized_question is None:
                    continue

                if normalized_question.startswith("="):
                    raise InvalidQuestionValueError(
                        f"worksheet {worksheet.title!r}, row {source_row}: "
                        "formula expressions are not accepted as question text"
                    )

                if len(normalized_question) > 4000:
                    raise InvalidQuestionValueError(
                        f"worksheet {worksheet.title!r}, row {source_row}: "
                        "question text exceeds 4000 characters"
                    )

                if len(questions) >= active_policy.max_questions:
                    raise TooManyQuestionsError(
                        f"workbook contains more than "
                        f"{active_policy.max_questions} questions"
                    )

                ordinal = len(questions) + 1

                question = build_question(
                    ordinal=ordinal,
                    sheet_name=worksheet.title,
                    source_row=source_row,
                    question_text=normalized_question,
                    section_path=current_section,
                )

                questions.append(question)
                sheet_question_count += 1

            imported_sheets.append(
                ImportedSheet(
                    sheet_name=worksheet.title,
                    header_row=header_row,
                    question_header=question_header_text,
                    section_header=section_header_text,
                    question_count=sheet_question_count,
                )
            )

        if not imported_sheets:
            raise NoQuestionColumnError(
                "no worksheet contains a supported question header"
            )

        if not questions:
            raise NoQuestionColumnError(
                "supported question headers were found, but no question rows "
                "contained usable question text"
            )

        questionnaire = build_questionnaire(
            name=questionnaire_name or _questionnaire_name_from_filename(filename),
            source_filename=filename,
            questions=tuple(questions),
        )

        return XlsxImportResult(
            questionnaire=questionnaire,
            imported_sheets=tuple(imported_sheets),
            ignored_sheets=tuple(ignored_sheets),
        )

    finally:
        workbook.close()


__all__ = [
    "XLSX_IMPORT_VERSION",
    "import_xlsx",
    "normalize_header",
]